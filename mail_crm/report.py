"""
report.py — экспорт CRM-данных в Excel.

Создаёт mail_crm_report.xlsx с листами:
  1. Контакты
  2. Треды активные
  3. Висяки — мне надо ответить
  4. Висяки — ждём от них
  5. Замороженные
  6. Сводка
"""
from __future__ import annotations

import logging
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import db

logger = logging.getLogger("report")

REPORT_PATH = Path(__file__).parent / "mail_crm_report.xlsx"

ACTIVE_DAYS = 14
FOLLOWUP_DAYS = 3

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_sheet(ws, header_count: int) -> None:
    """Форматирование: заголовки, авто-ширина, фильтры, фриз."""
    for col in range(1, header_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN

    # авто-ширина по содержимому (с лимитом 60)
    for col in range(1, header_count + 1):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter]:
            value = "" if cell.value is None else str(cell.value)
            for line in value.split("\n"):
                max_len = max(max_len, len(line))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 60)

    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def write_rows(ws, headers: list[str], rows: list[tuple]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_sheet(ws, len(headers))


def sheet_contacts(wb, conn) -> None:
    ws = wb.create_sheet("Контакты")
    rows = conn.execute("""
        SELECT email, name, company, sphere, relation_type,
               first_seen, last_seen, sent_count, received_count
        FROM contacts
        ORDER BY last_seen DESC
    """).fetchall()
    write_rows(
        ws,
        ["Email", "Имя", "Компания", "Сфера", "Тип",
         "Первый контакт", "Последний контакт", "Отправлено", "Получено"],
        [(r["email"], r["name"], r["company"], r["sphere"], r["relation_type"],
          r["first_seen"], r["last_seen"], r["sent_count"], r["received_count"])
         for r in rows],
    )


def _thread_rows(conn, where: str, params: tuple = ()) -> list[tuple]:
    sql = f"""
        SELECT subject_normalized, participants, last_message_date,
               days_since_last, awaiting_response_from, message_count
        FROM threads
        WHERE {where}
        ORDER BY days_since_last DESC
    """
    rows = conn.execute(sql, params).fetchall()
    out = []
    for r in rows:
        participants = ", ".join(db.json_loads(r["participants"]) or [])
        awaiting_label = {
            "me": "от меня",
            "them": "от них",
        }.get(r["awaiting_response_from"], "—")
        out.append((
            r["subject_normalized"] or "(без темы)",
            participants,
            r["last_message_date"],
            r["days_since_last"],
            awaiting_label,
            r["message_count"],
        ))
    return out


def sheet_active(wb, conn) -> None:
    ws = wb.create_sheet("Треды активные")
    rows = _thread_rows(conn, "days_since_last <= ?", (ACTIVE_DAYS,))
    write_rows(
        ws,
        ["Тема", "Участники", "Последнее письмо",
         "Дней назад", "Жду ответа от", "Кол-во писем"],
        rows,
    )


def sheet_my_turn(wb, conn) -> None:
    ws = wb.create_sheet("Висяки — мне надо ответить")
    rows = _thread_rows(conn, "awaiting_response_from = 'me'")
    write_rows(
        ws,
        ["Тема", "Участники", "Последнее письмо",
         "Дней назад", "Жду ответа от", "Кол-во писем"],
        rows,
    )


def sheet_their_turn(wb, conn) -> None:
    ws = wb.create_sheet("Висяки — ждём от них")
    rows = _thread_rows(
        conn,
        "awaiting_response_from = 'them' AND days_since_last >= ?",
        (FOLLOWUP_DAYS,),
    )
    write_rows(
        ws,
        ["Тема", "Участники", "Последнее письмо",
         "Дней назад", "Жду ответа от", "Кол-во писем"],
        rows,
    )


def sheet_frozen(wb, conn) -> None:
    """
    Замороженные: тред не активен (>14 дней), но среди писем треда есть
    offer/negotiation — потенциально упущенная сделка.
    """
    ws = wb.create_sheet("Замороженные")
    rows = conn.execute("""
        SELECT DISTINCT t.subject_normalized, t.participants,
                        t.last_message_date, t.days_since_last,
                        t.awaiting_response_from, t.message_count
        FROM threads t
        JOIN thread_messages tm ON tm.thread_id = t.id
        JOIN messages m ON m.message_id = tm.message_id
        WHERE t.days_since_last > ?
          AND m.category IN ('offer', 'negotiation')
        ORDER BY t.days_since_last DESC
    """, (ACTIVE_DAYS,)).fetchall()

    out = []
    for r in rows:
        participants = ", ".join(db.json_loads(r["participants"]) or [])
        awaiting_label = {
            "me": "от меня",
            "them": "от них",
        }.get(r["awaiting_response_from"], "—")
        out.append((
            r["subject_normalized"] or "(без темы)",
            participants,
            r["last_message_date"],
            r["days_since_last"],
            awaiting_label,
            r["message_count"],
        ))
    write_rows(
        ws,
        ["Тема", "Участники", "Последнее письмо",
         "Дней назад", "Жду ответа от", "Кол-во писем"],
        out,
    )


def sheet_summary(wb, conn) -> None:
    ws = wb.create_sheet("Сводка")

    total_msgs = conn.execute("SELECT COUNT(*) AS c FROM messages").fetchone()["c"]
    total_contacts = conn.execute("SELECT COUNT(*) AS c FROM contacts").fetchone()["c"]
    active_threads = conn.execute(
        "SELECT COUNT(*) AS c FROM threads WHERE days_since_last <= ?",
        (ACTIVE_DAYS,),
    ).fetchone()["c"]
    my_turn = conn.execute(
        "SELECT COUNT(*) AS c FROM threads WHERE awaiting_response_from = 'me'"
    ).fetchone()["c"]
    their_turn = conn.execute(
        "SELECT COUNT(*) AS c FROM threads "
        "WHERE awaiting_response_from = 'them' AND days_since_last >= ?",
        (FOLLOWUP_DAYS,),
    ).fetchone()["c"]

    ws.append(["Показатель", "Значение"])
    ws.append(["Всего писем", total_msgs])
    ws.append(["Всего контактов", total_contacts])
    ws.append([f"Активных тредов (≤ {ACTIVE_DAYS} дн.)", active_threads])
    ws.append(["Висяки — мне надо ответить", my_turn])
    ws.append([f"Висяки — ждём от них (≥ {FOLLOWUP_DAYS} дн.)", their_turn])

    ws.append([])
    ws.append(["ТОП-10 контактов по объёму переписки"])
    header_row = ws.max_row + 1
    ws.append(["Email", "Имя", "Компания",
               "Отправлено", "Получено", "Всего"])
    top = conn.execute("""
        SELECT email, name, company, sent_count, received_count,
               (sent_count + received_count) AS total
        FROM contacts
        ORDER BY total DESC
        LIMIT 10
    """).fetchall()
    for r in top:
        ws.append([r["email"], r["name"], r["company"],
                   r["sent_count"], r["received_count"], r["total"]])

    # Стилизуем шапку основной таблицы
    for col in range(1, 3):
        c = ws.cell(row=1, column=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = HEADER_ALIGN
    # И шапку ТОП-10
    for col in range(1, 7):
        c = ws.cell(row=header_row, column=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = HEADER_ALIGN

    # авто-ширина
    for col in range(1, 7):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 14), 60)

    ws.freeze_panes = "A2"


def main() -> int:
    db.setup_logging()
    db.init_schema()

    conn = db.get_conn()
    try:
        wb = Workbook()
        # Удаляем дефолтный пустой лист — будем создавать свои
        wb.remove(wb.active)

        sheet_contacts(wb, conn)
        sheet_active(wb, conn)
        sheet_my_turn(wb, conn)
        sheet_their_turn(wb, conn)
        sheet_frozen(wb, conn)
        sheet_summary(wb, conn)

        wb.save(REPORT_PATH)
        logger.info("Отчёт сохранён: %s", REPORT_PATH)
    finally:
        conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
